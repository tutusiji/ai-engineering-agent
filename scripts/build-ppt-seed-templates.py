#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build-ppt-seed-templates.py — 内置完整版式 PPT 模板打包器

把程序化绘制的页面背景（封面/章节/内容/结尾四张 1920×1080 jpg）组装为真实的
.pptx 模板文件：每套 4 页、每页一张近全幅图片，并用注入 clrScheme/fontScheme 的
theme1.xml 替换 python-pptx 默认 Office 主题——parseTemplate 据此提取配色/字体/
四类页面背景，buildPptx 按 pageType 铺对应底图。
生成内容全部为本程序原创绘制，无任何第三方素材，无版权负担。

输出: apps/studio-api/assets/ppt-seeds/tpl-<key>.pptx
用法: python3 scripts/build-ppt-seed-templates.py [资产目录]
"""

import io
import os
import re
import sys
import zipfile

from pptx import Presentation
from pptx.util import Emu

# 与 generate-ppt-seed-page-backgrounds.py 的 PALETTES 一致（六槽：primary/secondary/background/surface/text/accent）
PALETTES = {
    'luxe-indigo': {
        'primary': '4F46E5', 'secondary': '8B5CF6', 'background': 'FAF9F6',
        'surface': 'EEF2FF', 'text': '1E1B4B', 'accent': 'D4AF37',
    },
    'champagne-gold': {
        'primary': 'A16207', 'secondary': 'CA8A04', 'background': 'FDFBF6',
        'surface': 'FAF3E0', 'text': '292524', 'accent': '0E7490',
    },
    'jade-night': {
        'primary': '34D399', 'secondary': '2DD4BF', 'background': '0A0F0D',
        'surface': '12201A', 'text': 'E7F6EF', 'accent': 'FBBF24',
    },
    'mist-blue': {
        'primary': '0369A1', 'secondary': '38BDF8', 'background': 'FFFFFF',
        'surface': 'F0F9FF', 'text': '0C4A6E', 'accent': '059669',
    },
}

# 页序 → 页面背景语义（与 parse.ts 的多页提取映射一一对应）
PAGE_IMAGES = ['{key}.jpg', '{key}-section.jpg', '{key}-content.jpg', '{key}-ending.jpg']

# 16:9 标准画幅（EMU）
SLIDE_CX, SLIDE_CY = 12192000, 6858000

FONT = 'Microsoft YaHei'


def theme_xml(key, p):
    """构造注入用 theme1.xml 片段：clrScheme 六槽 + fontScheme（ YaHei）"""
    clr = (
        f'<a:clrScheme name="{key}">'
        f'<a:dk1><a:srgbClr val="{p["text"]}"/></a:dk1>'
        f'<a:lt1><a:srgbClr val="{p["background"]}"/></a:lt1>'
        f'<a:dk2><a:srgbClr val="{p["text"]}"/></a:dk2>'
        f'<a:lt2><a:srgbClr val="{p["surface"]}"/></a:lt2>'
        f'<a:accent1><a:srgbClr val="{p["primary"]}"/></a:accent1>'
        f'<a:accent2><a:srgbClr val="{p["secondary"]}"/></a:accent2>'
        f'<a:accent3><a:srgbClr val="{p["accent"]}"/></a:accent3>'
        f'<a:accent4><a:srgbClr val="{p["secondary"]}"/></a:accent4>'
        f'<a:accent5><a:srgbClr val="{p["primary"]}"/></a:accent5>'
        f'<a:accent6><a:srgbClr val="{p["accent"]}"/></a:accent6>'
        f'<a:hlink><a:srgbClr val="{p["primary"]}"/></a:hlink>'
        f'<a:folHlink><a:srgbClr val="{p["secondary"]}"/></a:folHlink>'
        f'</a:clrScheme>'
    )
    fonts = (
        f'<a:fontScheme name="{key}">'
        f'<a:majorFont><a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/><a:cs typeface=""/></a:majorFont>'
        f'<a:minorFont><a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/><a:cs typeface=""/></a:minorFont>'
        f'</a:fontScheme>'
    )
    return clr, fonts


def inject_theme(pptx_bytes, key, p):
    """zip 级替换 theme1.xml 的 clrScheme/fontScheme（python-pptx 不暴露主题 API）"""
    clr, fonts = theme_xml(key, p)
    src = zipfile.ZipFile(io.BytesIO(pptx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == 'ppt/theme/theme1.xml':
                xml = data.decode('utf-8')
                xml, n1 = re.subn(r'<a:clrScheme[^>]*>.*?</a:clrScheme>', clr, xml, flags=re.S)
                xml, n2 = re.subn(r'<a:fontScheme[^>]*>.*?</a:fontScheme>', fonts, xml, flags=re.S)
                if n1 != 1 or n2 != 1:
                    raise RuntimeError(f'{key}: theme1.xml 替换失败 clrScheme={n1} fontScheme={n2}')
                data = xml.encode('utf-8')
            dst.writestr(item, data)
    return out.getvalue()


def build_template(asset_dir, key):
    """组装单套模板：4 页整页背景图 + 注入主题色/字体，返回 .pptx 字节"""
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_CX)
    prs.slide_height = Emu(SLIDE_CY)
    for pattern in PAGE_IMAGES:
        img_path = os.path.join(asset_dir, pattern.format(key=key))
        if not os.path.isfile(img_path):
            raise FileNotFoundError(f'页面背景缺失: {img_path}')
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank 版式（无占位符）
        slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
    buf = io.BytesIO()
    prs.save(buf)
    return inject_theme(buf.getvalue(), key, PALETTES[key])


def main():
    asset_dir = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), '..', 'apps', 'studio-api', 'assets', 'ppt-seeds')
    if not os.path.isdir(asset_dir):
        raise SystemExit(f'资产目录不存在: {asset_dir}')
    for key in PALETTES:
        out_path = os.path.join(asset_dir, f'tpl-{key}.pptx')
        data = build_template(asset_dir, key)
        with open(out_path, 'wb') as f:
            f.write(data)
        print(f'tpl-{key}.pptx: {len(data) // 1024} KB（4 页）')
    print('完成：4 套完整版式模板')


if __name__ == '__main__':
    main()
